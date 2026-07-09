"""
DOCX document parser.

Uses python-docx to extract text from Microsoft Word documents.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.logging_config import get_logger
from app.ingestion.parsers.base_parser import BaseParser

logger = get_logger(__name__)


class DocxParser(BaseParser):
    """
    Parser for Microsoft Word (.docx) documents.
    """

    supported_extensions = (".docx",)

    def parse(self, file_path: Path) -> str:
        """
        Parse a DOCX document.

        Args:
            file_path: Path to the DOCX document.

        Returns:
            Extracted document text.

        Raises:
            ValueError:
                If the document cannot be parsed.
        """
        self.validate(file_path)

        logger.info("Parsing DOCX document: %s", file_path.name)

        try:
            document = Document(file_path)

            paragraphs: list[str] = []

            # Extract paragraphs
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text:
                    paragraphs.append(text)

            # Extract table data
            for table in document.tables:
                for row in table.rows:
                    row_data = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text.strip()
                    ]

                    if row_data:
                        paragraphs.append(" | ".join(row_data))

            extracted_text = "\n".join(paragraphs).strip()

            logger.info(
                "Successfully extracted %d characters from '%s'.",
                len(extracted_text),
                file_path.name,
            )

            return extracted_text

        except Exception as exc:
            logger.exception(
                "Failed to parse DOCX '%s'.",
                file_path.name,
            )
            raise ValueError(
                f"Unable to parse DOCX '{file_path.name}'."
            ) from exc